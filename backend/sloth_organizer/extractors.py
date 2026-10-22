"""
Content preview extraction for supported file formats.
Previews are capped at MAX_PREVIEW_CHARS to control AI token usage.
"""

from pathlib import Path

MAX_PREVIEW_CHARS = 500


def extract_preview(path: Path, extension: str) -> str | None:
    """Return a short text preview of *path*, or None if unsupported / unreadable."""
    ext = extension.lower()
    try:
        if ext in {".txt", ".md", ".csv", ".json"}:
            return _read_text(path)
        if ext == ".pdf":
            return _read_pdf(path)
        if ext == ".docx":
            return _read_docx(path)
        if ext == ".xlsx":
            return _read_xlsx(path)
    except Exception:
        # Never crash the scan because of a bad file
        pass
    return None


# ---------------------------------------------------------------------------
# Plain-text formats
# ---------------------------------------------------------------------------

def _read_text(path: Path) -> str | None:
    try:
        with path.open("r", encoding="utf-8", errors="replace") as fh:
            return fh.read(MAX_PREVIEW_CHARS).strip() or None
    except OSError:
        return None


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _read_pdf(path: Path) -> str | None:
    from pypdf import PdfReader  # lazy import — only when needed
    reader = PdfReader(str(path))
    if not reader.pages:
        return None
    text = reader.pages[0].extract_text() or ""
    return text[:MAX_PREVIEW_CHARS].strip() or None


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def _read_docx(path: Path) -> str | None:
    from docx import Document  # lazy import
    doc = Document(str(path))
    parts: list[str] = []
    total = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        parts.append(text)
        total += len(text)
        if total >= MAX_PREVIEW_CHARS:
            break
    return " ".join(parts)[:MAX_PREVIEW_CHARS].strip() or None


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def _read_xlsx(path: Path) -> str | None:
    from openpyxl import load_workbook  # lazy import
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        return None
    cells: list[str] = []
    total = 0
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell is None:
                continue
            text = str(cell).strip()
            if not text:
                continue
            cells.append(text)
            total += len(text)
            if total >= MAX_PREVIEW_CHARS:
                break
        if total >= MAX_PREVIEW_CHARS:
            break
    wb.close()
    return " | ".join(cells)[:MAX_PREVIEW_CHARS].strip() or None
